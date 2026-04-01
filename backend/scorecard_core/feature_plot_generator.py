import pandas as pd
import numpy as np
import io
import os
import sys

from matplotlib import pyplot as plt
from openpyxl.drawing.image import Image
from tqdm import tqdm
# sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.dirname(__file__))))
# sys.path.append("D:/work/scorecard")
from . import decisiontree_bins

def create_feature_plot_report(excel_writer, data, feature_list, sheet_name='特征分箱图',if_all=False):
    """
    为给定的特征列表自动生成特征分箱图报告，并将其添加到一个已存在的ExcelWriter对象中。

    该函数会执行以下操作:
    1. 在训练集上计算每个特征的最佳分箱和IV值。
    2. 在指定的ExcelWriter中创建一个新的工作表。
    3. 在工作表中写入按IV值排序的特征名及其IV值。
    4. 使用训练集的分箱标准，为每个特征在 'train', 'valid', 'oot' 数据集上生成分箱图。
    5. 将生成的图表插入到Excel工作表的相应位置。

    参数:
    - excel_writer (pd.ExcelWriter): 一个Pandas ExcelWriter对象，报告将被写入其中。
    - data (pd.DataFrame): 包含特征列、'label'列和'target'列的源数据。
                           'target'列应包含 'train', 'valid', 'oot' 等值。
    - feature_list (list): 需要进行分析和绘图的特征名称列表。
    - sheet_name (str): 在Excel中创建的工作表的名称。
    """

    print(f"--- 正在向工作表 '{sheet_name}' 添加特征分箱图 ---")
    
    # 检查输入数据
    if 'label' not in data.columns or 'target' not in data.columns:
        raise ValueError("输入数据 'data' 中必须包含 'label' 和 'target' 列。")

    # 1. 在训练集上计算IV和最优分箱
    print("Step 1/4: 在训练集上计算特征IV值和分箱...")
    feature_metrics = []
    train_data = data[data['target'] == 'train'].copy()
    if train_data.empty:
        raise ValueError("数据中找不到 'target' 为 'train' 的数据，无法计算基准IV和分箱。")

    for feature in tqdm(feature_list, desc="计算IV和分箱"):
        if feature not in train_data.columns:
            print(f"  [警告] 特征 '{feature}' 在训练集中不存在，已跳过。")
            continue
        try:
            # 在训练集上计算最优分箱边界
            boundary = decisiontree_bins.optimal_binning_boundary(train_data[feature], train_data['label'])
            # 使用此边界计算IV
            bin_result = decisiontree_bins.feature_woe_iv_bins(train_data[feature], train_data['label'], bins=boundary)
            iv = bin_result['iv'].sum()
            feature_metrics.append({'feature': feature, 'iv': iv, 'bins': boundary})
        except Exception as e:
            print(f"  [错误] 计算特征 '{feature}' 的IV值失败: {e}")
            feature_metrics.append({'feature': feature, 'iv': -1, 'bins': None}) # 标记为错误

    # 根据IV值降序排序
    sorted_features = sorted(feature_metrics, key=lambda x: x['iv'], reverse=True)

    # 2. 在传入的excel_writer中创建工作表
    print(f"Step 2/4: 在Excel中创建工作表 '{sheet_name}'...")
    pd.DataFrame().to_excel(excel_writer, sheet_name=sheet_name, index=False)
    wb = excel_writer.book
    ws = wb[sheet_name]

    # 3. 写入表头和排序后的特征信息
    print("Step 3/4: 写入特征信息...")
    ws['A1'] = '特征名称'
    if if_all:
        ws['B1'] = 'IV值'
    else:
        ws['B1'] = 'IV值 (基于训练集)'
        
    IMAGE_COLUMNS = {'train': 'C', 'valid': 'J', 'oot': 'P'}

    if if_all:
        IMAGE_COLUMNS_all = {'total': 'C'}
        for target, col_letter in IMAGE_COLUMNS_all.items():
            ws.column_dimensions[col_letter].width = 70
            ws[f'{col_letter}1'] = f"{target.upper()} 分箱图"
    else:
        for target, col_letter in IMAGE_COLUMNS.items():
            ws.column_dimensions[col_letter].width = 70
            ws[f'{col_letter}1'] = f"{target.upper()} 分箱图"
    ws.column_dimensions['A'].width = 20
    ws.column_dimensions['B'].width = 20

    for idx, item in enumerate(sorted_features):
        row = idx + 2
        ws[f'A{row}'] = item['feature']
        ws[f'B{row}'] = round(item['iv'], 6) if item['iv'] != -1 else '计算失败'

    # 4. 遍历特征并生成图表
    print("Step 4/4: 生成并插入特征分箱图...")
    for idx, item in enumerate(tqdm(sorted_features, desc="生成图表")):
        excel_row_index = idx + 2
        var_name = item['feature']
        bins = item['bins']
        
        if bins is None:
            print(f"  [跳过] 特征 '{var_name}' 因无有效分箱而跳过绘图。")
            continue

        ws.row_dimensions[excel_row_index].height = 330

        for target, col_letter in IMAGE_COLUMNS.items():
            temp_df = data[data['target'] == target].copy()
            
            if temp_df.empty or var_name not in temp_df.columns:
                continue
            
            try:
                result = decisiontree_bins.feature_woe_iv_bins(temp_df[var_name], temp_df['label'], bins=bins)
                
                if if_all:
                    fig = decisiontree_bins.plot_bin(result, title=f"{var_name} (TOTAL)")
                else:
                    fig = decisiontree_bins.plot_bin(result, title=f"{var_name} ({target.upper()})")
                
                img_io = io.BytesIO()
                fig.savefig(img_io, format='png', bbox_inches='tight')
                img_io.seek(0)
                img = Image(img_io)
                
                cell_address = f"{col_letter}{excel_row_index}"
                ws.add_image(img, cell_address)
                
                plt.close(fig)
            except Exception as e:
                print(f"    [错误] 为变量 '{var_name}' 在 '{target}' 上生成图表失败: {e}")
                cell_address = f"{col_letter}{excel_row_index}"
                ws[f'{cell_address}'] = f"该特征分布原因，图表生成失败"

    print(f"--- 特征分箱图已成功添加至工作表 '{sheet_name}' ---")


if __name__ == '__main__':
    # --- 这是一个使用示例 ---
    # 在实际使用中，请注释或删除这部分代码，并从外部调用 create_feature_plot_report 函数

    # 1. 创建一个模拟的 scored_data DataFrame
    print("\n--- 开始运行使用示例 ---")
    print("正在创建模拟数据...")
    num_samples = 1000
    features = {
        'credit_score': np.random.randint(300, 850, num_samples),
        'loan_amount': np.random.rand(num_samples) * 50000 + 5000,
        'annual_income': np.random.randn(num_samples) * 30000 + 90000,
        'months_since_last_delinquency': np.random.randint(0, 100, num_samples)
    }
    df_data = pd.DataFrame(features)
    df_data['label'] = np.random.choice([0, 1], num_samples, p=[0.9, 0.1])
    
    # 创建 'target' 列来区分 train, valid, oot
    split = np.random.choice(['train', 'valid', 'oot'], num_samples, p=[0.7, 0.15, 0.15])
    df_data['target'] = split
    print("模拟数据创建完成。")

    # 2. 定义你要画图的特征列表
    features_to_plot = list(features.keys())

    # 3. 指定输出路径和工作表名
    output_excel_path = 'sample_feature_report_integrated.xlsx'
    report_sheet_name = '特征分箱报告'
    
    # 4. 调用函数生成报告 (集成模式)
    print(f"正在使用集成模式生成报告到 '{output_excel_path}'...")
    with pd.ExcelWriter(output_excel_path, engine='openpyxl') as writer:
        # 假设这里先写入了其他数据
        pd.DataFrame({'info': ['some other analysis']}).to_excel(writer, sheet_name='概要', index=False)

        # 调用我们的绘图函数
        create_feature_plot_report(
            excel_writer=writer,
            data=df_data,
            feature_list=features_to_plot,
            sheet_name=report_sheet_name
        )

    print(f"\n示例报告已生成: '{output_excel_path}'")
    print("--- 示例运行结束 ---")
