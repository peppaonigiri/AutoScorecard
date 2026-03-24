import pandas as pd
import warnings
import re
from pathlib import Path
from openpyxl.styles import Alignment

warnings.filterwarnings('ignore')

# 预编译正则
DATE_PATTERNS = [
    re.compile(r"(\d{4})[-/\.](\d{1,2})"),
    re.compile(r"(\d{4})(\d{2})"),
    re.compile(r"(\d{4})年(\d{1,2})月")
]

class UniversalReportGenerator:
    """三合一模型评估报告生成器"""
    def __init__(self, base_path):
        self.base_path = Path(base_path)
        self.df_raw = pd.DataFrame() 
        self.pivot_tables = {}

    @staticmethod
    def extract_time(text):
        """从文本提取时间范围"""
        text = str(text)
        found = set()
        for pat in DATE_PATTERNS:
            for y, m in pat.findall(text):
                try:
                    if 1900 <= int(y) <= 2100 and 1 <= int(m) <= 12:
                        found.add(f"{int(y)}-{int(m):02d}")
                except ValueError:
                    continue
        
        found = sorted(found)
        if not found:
            return "未知"
        return f"{found[0]} to {found[-1]}" if len(found) >= 2 else found[0]

    def parse_model_name(self, name):
        """解析模型名称"""
        parts = name.split('_')
        if not parts:
            return "未知", "未知", "未知"
            
        data_source = parts[0]
        if len(parts) >= 3:
            return data_source, parts[1], parts[2]
        
        # 处理只有两部分的情况
        part2 = parts[1] if len(parts) > 1 else "未知"
        if "target" in part2:
            return data_source, "未知", part2
        return data_source, part2, "未知"

    def get_excel_info(self, path):
        """提取Excel中的数据概要和时间"""
        try:
            df = pd.read_excel(path, sheet_name="数据概要", header=None)
        except Exception:
            return {}, ("未知", "未知")

        info = {}
        times = {'oot': [], 'tv': [], 'all': []}

        for _, row in df.iterrows():
            row_list = list(row.dropna())
            if not row_list:
                continue

            # 提取统计信息 (count, badrate, etc.)
            if len(row) >= 3 and pd.notna(row[0]):
                key = str(row[0]).lower()
                info[key] = {
                    'count': row_list[1] if len(row_list) > 1 else "未知",
                    'badrate': row_list[2] if len(row_list) > 2 else "未知",
                    'badrate_weight': row_list[3] if len(row_list) > 3 else "未知"
                }
            
            # 提取时间信息
            txt = " ".join(map(str, row_list))
            t = self.extract_time(txt)
            if t != "未知":
                times['all'].append(t)
                txt_lower = txt.lower()
                if 'oot' in txt_lower:
                    times['oot'].append(t)
                if any(x in txt_lower for x in ['train', 'valid']):
                    times['tv'].append(t)

        def merge_dates(ts): 
            if not ts:
                return "未知"
            # 拆分所有时间段并重新排序
            dates = set()
            for t in ts:
                parts = t.split(' to ')
                dates.update(parts)
            dates = sorted(dates)
            return f"{dates[0]} to {dates[-1]}" if len(dates) >= 2 else dates[0]

        return info, (merge_dates(times['oot']), merge_dates(times['tv']))

    def run_aggregation(self, ignore_folders=None):
        """汇总基础数据"""
        print(f"[INFO] 正在扫描路径: {self.base_path} ...")
        files = list(self.base_path.rglob("file/建模报告/模型评估报告.xlsx"))
        if ignore_folders:
            files = [f for f in files if not set(f.parts) & set(ignore_folders)]
        
        if not files:
            print("[WARN] 未找到任何评估报告文件")
            return

        dfs = []
        for f in files:
            try:
                df = pd.read_excel(f, sheet_name="模型性能")
                
                # 解析路径获取模型名
                # 假设结构为 .../ModelName/file/建模报告/...
                parts = f.parts
                try:
                    idx = parts.index("file")
                    model_name = parts[idx-1] if idx > 0 else "未知模型"
                except ValueError:
                    model_name = "未知模型"
                
                ds, cg, target_var = self.parse_model_name(model_name)
                d_info, (t_oot, t_tv) = self.get_excel_info(f)
                
                # 标准化数据集名称
                df['datasets'] = df['datasets'].astype(str).str.lower().apply(
                    lambda x: next((k for k in ['train', 'valid', 'oot'] if k in x), x)
                )

                # 填充字段
                df['model_name'] = model_name
                df['data_source'] = ds
                df['customer_group'] = cg
                df['target_var'] = target_var
                
                # 时间范围映射
                df['time_range'] = df['datasets']
                mask_oot = df['datasets'] == 'oot'
                mask_tv = df['datasets'].isin(['train', 'valid'])
                
                df.loc[mask_oot, 'time_range'] = t_oot
                df.loc[mask_tv, 'time_range'] = t_tv
                
                # 统计指标映射
                for col in ['count', 'badrate', 'badrate_weight']:
                    df[col] = df['datasets'].apply(lambda d: d_info.get(d, {}).get(col, '未知'))
                
                dfs.append(df)
            except Exception as e:
                print(f"[ERROR] 读取异常 {f}: {e}")

        if dfs:
            self.df_raw = pd.concat(dfs, ignore_index=True)
            self._post_process_raw_data()
            print(f"[INFO] 基础汇总完成，共 {len(self.df_raw)} 行")

    def _post_process_raw_data(self):
        """处理原始数据的排序和清理"""
        # 排序
        dataset_order = {'train': 1, 'valid': 2, 'oot': 3}
        self.df_raw['sort_key'] = self.df_raw['datasets'].map(lambda x: dataset_order.get(x, 4))
        self.df_raw = self.df_raw.sort_values(['model_name', 'sort_key', 'datasets']).drop(columns=['sort_key'])
        
        # 插入序号
        self.df_raw.insert(0, 'model_index', self.df_raw.groupby('model_name', sort=False).ngroup() + 1)
        
        # 清理无效列
        for col in ['target_var', 'badrate_weight']:
            if col in self.df_raw and self.df_raw[col].isin(['未知', None]).all():
                self.df_raw.drop(columns=[col], inplace=True)

    def _sort_dataframe(self, df, col_order, row_order):
        """统一辅助排序方法"""
        if 'data_source' in df.columns and row_order:
            df['data_source'] = pd.Categorical(df['data_source'], categories=self._get_sorted_categories(df['data_source'], row_order), ordered=True)
        
        if 'datasets' in df.columns:
            # Bug fix: use _get_sorted_categories to include values like months that are not in the predefined list
            df['datasets'] = pd.Categorical(df['datasets'], categories=self._get_sorted_categories(df['datasets'], ['train', 'valid', 'oot']), ordered=True)
            
        if 'customer_group' in df.columns and col_order:
            df['customer_group'] = pd.Categorical(df['customer_group'], categories=self._get_sorted_categories(df['customer_group'], col_order), ordered=True)
        
        return df

    def _get_sorted_categories(self, series, simplify_order):
        """获取排序后的分类列表"""
        unique_vals = series.dropna().unique()
        priority = simplify_order if simplify_order else []
        final = [x for x in priority if x in unique_vals] + sorted([x for x in unique_vals if x not in priority])
        return final

    def run_pivot(self, configs, col_order, row_order, swap=False):
        """生成透视表"""
        if self.df_raw.empty:
            return

        df = self.df_raw.copy()
        self._sort_dataframe(df, col_order, row_order)

        idx = ['customer_group', 'datasets'] if swap else ['data_source', 'datasets']
        cols = ['data_source'] if swap else ['customer_group']
        if 'target_var' in df:
            cols.append('target_var')

        for cfg in configs:
            val_col = cfg['value_column']
            if val_col not in df:
                continue
            
            # 处理 '未知'
            if df[val_col].isin(['未知']).any():
                df[val_col] = df[val_col].replace('未知', None)

            pt = df.pivot_table(
                index=idx, 
                columns=cols, 
                values=val_col, 
                aggfunc=cfg.get('aggfunc', 'mean'), 
                fill_value=cfg.get('fill_value', "/")
            )
            
            # 移除全空行
            mask = pt.apply(lambda r: any(x not in [cfg.get('fill_value', "/"), 0, None] for x in r), axis=1)
            if not pt[mask].empty:
                self.pivot_tables[cfg['name']] = pt[mask]
                print(f"  -> 生成透视表: {cfg['name']}")

        self._make_summary(configs, row_order if swap else col_order)

    def _make_summary(self, cfgs, sort_order):
        """生成 ALL_SUMMARY"""
        valid_names = [c['name'] for c in cfgs if c['name'] in self.pivot_tables]
        if not valid_names:
            return

        common_cols = {}
        unique_cols = {}
        base_df = self.pivot_tables[valid_names[0]]
        
        for name in valid_names:
            tbl = self.pivot_tables[name]
            # 检查是否为公共列 (所有列值相同)
            if tbl.shape[1] > 1 and all(tbl.iloc[:, 0].equals(tbl.iloc[:, i]) for i in range(1, tbl.shape[1])):
                common_cols[(' ', name)] = tbl.iloc[:, 0]
            else:
                for c in tbl.columns:
                    col_str = "_".join(map(str, c)) if isinstance(c, tuple) else str(c)
                    unique_cols[(col_str, name)] = tbl[c]

        if not common_cols and not unique_cols:
            return
        
        df_sum = pd.DataFrame({**common_cols, **unique_cols}, index=base_df.index)
        
        # 构建列顺序
        # 1. 公共列
        final_cols = sorted(common_cols.keys(), key=lambda x: valid_names.index(x[1]))
        
        # 2. 分组列
        groups = sorted(set(k[0] for k in unique_cols.keys()))
        
        # 自定义排序逻辑
        def get_group_sort_index(group_name):
            if not sort_order:
                return 999
            for i, order_val in enumerate(sort_order):
                if str(order_val) in group_name:
                    return i
            return 999

        groups.sort(key=get_group_sort_index)
            
        for g in groups:
            final_cols.extend([(g, m) for m in valid_names if (g, m) in unique_cols])
            
        df_sum = df_sum.reindex(columns=pd.MultiIndex.from_tuples(final_cols))
        self.pivot_tables = {'ALL_SUMMARY': df_sum, **self.pivot_tables}
        print(f"[INFO] ALL_SUMMARY 生成完毕")

    def run_psi(self, sources, cust_types, swap=False):
        """聚合PSI"""
        print("[INFO] 正在聚合 PSI ...")
        if self.df_raw.empty:
            return

        has_target = 'target_var' in self.df_raw.columns
        
        # 筛选有效模型
        mask = (self.df_raw['data_source'].isin(sources)) & (self.df_raw['customer_group'].isin(cust_types))
        cols_needed = ['model_name', 'data_source', 'customer_group']
        if has_target:
            cols_needed.append('target_var')
        
        valid_models = self.df_raw.loc[mask, cols_needed].drop_duplicates()
        
        frames = []
        for _, row in valid_models.iterrows():
            m_name = row['model_name']
            path = self.base_path / m_name / 'file/建模报告/模型评估报告.xlsx'
            
            if not path.exists():
                continue
            
            try:
                frames.append(self._read_psi_sheet(path, row))
            except Exception as e:
                print(f"[ERROR] 读取PSI异常 {m_name}: {e}")
        
        frames = [f for f in frames if f is not None]
        if not frames:
            print("[WARN] 无PSI数据")
            return
        
        df_long = pd.concat(frames)
        
        # 生成透视表
        idx = ['customer_group', 'datasets'] if swap else ['data_source', 'datasets']
        cols = ['data_source'] if swap else ['customer_group']
        if has_target:
            cols.append('target_var')

        try:
            out = df_long.pivot_table(index=idx, columns=cols, values='psi', aggfunc='first')
            # 排序
            sort_keys = sources if swap else cust_types
            out = self._sort_pivot_columns(out, sort_keys)
            
            self.pivot_tables['PSI'] = out
            print("[INFO] PSI 聚合完成")
        except Exception as e:
            print(f"[ERROR] PSI 透视失败: {e}")

    def _read_psi_sheet(self, path, model_info):
        """读取单个文件的PSI Sheet"""
        xl = pd.ExcelFile(path)
        candidates = [s for s in xl.sheet_names if 'PSI' in s.upper()]
        if not candidates and len(xl.sheet_names) > 6:
            candidates = [xl.sheet_names[6]]
            
        for s_name in candidates:
            try:
                df = pd.read_excel(path, sheet_name=s_name)
                df.columns = df.columns.astype(str).str.lower()
                
                c_psi = next((c for c in df.columns if 'psi' in c), None)
                c_moth = next((c for c in df.columns if 'moth' in c or 'month' in c or 'date' in c), None)
                
                if c_psi and c_moth:
                    sub = df[[c_moth, c_psi]].rename(columns={c_psi: 'psi', c_moth: 'datasets'})
                    sub['data_source'] = model_info['data_source']
                    sub['customer_group'] = model_info['customer_group']
                    if 'target_var' in model_info:
                        sub['target_var'] = model_info['target_var']
                    return sub
            except:
                continue
        return None

    def _sort_pivot_columns(self, df, sort_keys):
        """对透视表列进行自定义排序"""
        if isinstance(df.columns, pd.MultiIndex):
            def col_key(col_tuple):
                main_key = col_tuple[0]
                try:
                    return sort_keys.index(main_key)
                except ValueError:
                    return len(sort_keys)
            new_cols = sorted(df.columns, key=col_key)
            return df.reindex(columns=new_cols)
        else:
            ordered = [c for c in sort_keys if c in df.columns]
            others = [c for c in df.columns if c not in ordered]
            return df.reindex(columns=ordered + others)

    def save(self, path):
        """保存结果"""
        print(f"[INFO] 保存至: {path}")
        with pd.ExcelWriter(path, engine='openpyxl') as writer:
            for n, t in self.pivot_tables.items():
                if not t.empty:
                    t.to_excel(writer, sheet_name=n)
            
            if not self.df_raw.empty:
                self._save_summary_sheet(writer)
        print("[INFO] 完成")

    def _save_summary_sheet(self, writer):
        """保存汇总数据Sheet并进行格式化"""
        cols = ['model_index', 'model_name', 'data_source', 'customer_group', 'target_var', 
                'datasets', 'time_range', 'count', 'badrate', 'badrate_weight', 'auc', 'ks', 'sub_5_lift', 'top_5_lift']
        
        valid_cols = [c for c in cols if c in self.df_raw]
        self.df_raw[valid_cols].to_excel(writer, sheet_name='汇总数据', index=False)
        
        ws = writer.sheets['汇总数据']
        align = Alignment(horizontal='center', vertical='center')
        
        # 寻找需要合并的列索引
        header_map = {cell.value: i+1 for i, cell in enumerate(ws[1])}
        merge_fields = ['model_index', 'model_name', 'data_source', 'customer_group', 'target_var']
        merge_indices = [header_map[f] for f in merge_fields if f in header_map]
        
        curr_row = 2
        for _, g in self.df_raw.groupby('model_name', sort=False):
            group_len = len(g)
            if group_len > 1:
                for idx in merge_indices:
                    ws.merge_cells(start_row=curr_row, start_column=idx, end_row=curr_row+group_len-1, end_column=idx)
            curr_row += group_len
            
        for row in ws.iter_rows():
            for cell in row:
                cell.alignment = align

if __name__ == "__main__":
    BASE = r"d:\work\06.std_product\wxjk\test_data_gen"
    OUT = r"d:\work\06.std_product\wxjk\test_data_gen\test_data_report.xlsx"
    SOURCES = ['delta', 'theta', 'zeta', 'ronghe']
    TYPES = ['xd', 'fd', 'xdjc','fdjc','xdfd','fdapp1','fdapp0']
    CONFIGS = [
        {'name': 'count', 'value_column': 'count'},
        {'name': 'badrate', 'value_column': 'badrate'},
        {'name': 'badrate_weight', 'value_column': 'badrate_weight'},
        {'name': 'auc', 'value_column': 'auc'},
        {'name': 'ks', 'value_column': 'ks'},
    ]
    
    gen = UniversalReportGenerator(BASE)
    gen.run_aggregation()
    gen.run_pivot(CONFIGS, col_order=TYPES, row_order=SOURCES, swap=True)
    gen.run_psi(SOURCES, TYPES, swap=True)
    gen.save(OUT)
